"""Origin 2021 task runner. Natural-language planning belongs to Codex; this runs reviewed JSON specs."""
from __future__ import annotations
import argparse, contextlib, datetime as dt, hashlib, importlib.metadata as meta
import json, math, msvcrt, os, pathlib, re, struct, subprocess, sys, time, traceback
P = pathlib.Path

def read_json(path):
    return json.loads(P(path).read_text(encoding="utf-8-sig"))

def write_json(path, data):
    P(path).parent.mkdir(parents=True, exist_ok=True)
    P(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")

def require(ok, message):
    if not ok:
        raise ValueError(message)

def keys(obj, allowed, label):
    require(isinstance(obj, dict), label + " must be an object")
    extra = set(obj) - set(allowed.split())
    require(not extra, f"{label}: unsupported fields {sorted(extra)}; use a reviewed custom Origin script, do not discard requested features.")

def sha(path):
    h = hashlib.sha256()
    with P(path).open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()

def origin_processes():
    result = subprocess.run(["powershell.exe", "-NoProfile", "-Command",
        "Get-Process Origin64 -ErrorAction SilentlyContinue | Select-Object Id,MainWindowTitle | ConvertTo-Json -Compress"],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=20)
    text = result.stdout.strip()
    data = json.loads(text) if text else []
    return data if isinstance(data, list) else [data]

def environment(config):
    require(os.name == "nt", "Origin automation requires Windows.")
    require(sys.version_info[:2] == (3, 11) and struct.calcsize("P") == 8,
            "Use the configured 64-bit Python 3.11, never the default Python.")
    require(P(sys.executable).resolve() == P(config["ORIGIN_PYTHON"]).resolve(),
            "Interpreter differs from ORIGIN_PYTHON in origin_config.json.")
    return {"python": sys.version, "python_executable": sys.executable,
            "packages": {n: meta.version(n) for n in ("originpro", "OriginExt", "pandas", "openpyxl", "numpy")}}

def read_table(source, sheet=0, header=0, delimiter=None, encoding="utf-8-sig"):
    import pandas as pd
    path = P(source).resolve()
    require(path.is_file(), f"Data file not found: {path}")
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        if sheet is None:
            sheet = 0
        # A formula without an Excel-calculated cache is not a missing experimental value.
        from openpyxl import load_workbook
        formulas=load_workbook(path,read_only=True,data_only=False)
        cached=load_workbook(path,read_only=True,data_only=True)
        try:
            name=formulas.sheetnames[sheet] if isinstance(sheet,int) else sheet
            require(name in formulas.sheetnames,"Requested Excel worksheet not found.")
            uncached=[]
            for frow,crow in zip(formulas[name].iter_rows(),cached[name].iter_rows()):
                for fc,cc in zip(frow,crow):
                    if fc.data_type=="f" and cc.value is None:
                        uncached.append(fc.coordinate)
                        if len(uncached)>=20: break
                if len(uncached)>=20: break
            require(not uncached, "Excel formulas have no cached values at "+str(uncached)+
                    ". Recalculate a COPY using Excel/LibreOffice or supply evaluated data; do not treat them as missing.")
        finally:
            formulas.close(); cached.close()
        df = pd.read_excel(path, sheet_name=sheet, header=header, engine="openpyxl")
    elif ext == ".xls":
        raise ValueError("Legacy .xls needs an explicitly verified reader or a local .xlsx conversion; no implicit package changes.")
    elif ext in (".csv", ".tsv", ".txt", ".dat"):
        sep = delimiter if delimiter is not None else ("\t" if ext == ".tsv" else "," if ext == ".csv" else None)
        df = pd.read_csv(path, sep=sep, header=header, encoding=encoding, engine="python", skip_blank_lines=False)
    else:
        raise ValueError(f"Unsupported table extension: {ext}")
    require(not df.empty, "Input table is empty.")
    df.columns = [str(c) for c in df.columns]
    require(len(set(df.columns)) == len(df.columns), "Column labels must be unique.")
    return df, path

def inspect_data(path, sheet):
    import pandas as pd
    df, path = read_table(path, sheet)
    report = {"source": str(path), "sha256": sha(path), "rows": len(df),
              "columns": [{"position_1based": i+1, "name": str(c), "dtype": str(df[c].dtype),
                           "missing": int(df[c].isna().sum())} for i,c in enumerate(df.columns)],
              "head": json.loads(df.head(8).to_json(orient="records", force_ascii=False))}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        with pd.ExcelFile(path, engine="openpyxl") as book:
            report["worksheets"] = book.sheet_names
    return report

def validate(spec, workspace):
    keys(spec, "schema_version id mode executor runtime data page layers plots axes annotations style export analysis notes", "FigureSpec")
    require(spec.get("schema_version") == 1, "schema_version must be 1.")
    require(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}", spec.get("id","")), "Use a short ASCII figure id.")
    require(spec.get("mode") in ("preview","execute"), "mode must explicitly be preview or execute.")
    require(spec.get("executor","core") == "core", "This spec needs a custom Origin 2021 script; read the skill recipes.")
    keys(spec.get("runtime",{}), "show_origin show_seconds", "runtime")
    require(0 <= float(spec.get("runtime",{}).get("show_seconds",3)) <= 60, "show_seconds must be 0..60.")
    require(isinstance(spec.get("data"),list) and spec["data"], "Provide data sources.")
    require(isinstance(spec.get("layers"),list) and spec["layers"], "Provide layers.")
    require(isinstance(spec.get("plots"),list) and spec["plots"], "Provide plots.")
    keys(spec.get("page",{}), "template", "page")
    keys(spec.get("style",{}), "font font_size line_width symbol_size colors fit_line_color hide_legend reference_calibration slope_unit", "style")
    style = spec.get("style",{})
    require(re.fullmatch(r"[A-Za-z0-9 -]+", style.get("font","Arial")), "Invalid font name.")
    for field,default in (("font_size",16),("line_width",1.5),("symbol_size",6)):
        require(0 < float(style.get(field,default)) <= 80, f"Invalid {field}.")
    for color in style.get("colors",["#2463A6"]):
        require(re.fullmatch(r"#[0-9a-fA-F]{6}",color), "Colors must be #RRGGBB.")
    if "fit_line_color" in style:
        require(re.fullmatch(r"#[0-9a-fA-F]{6}",style["fit_line_color"]), "Fit line color must be #RRGGBB.")
    if "hide_legend" in style:
        require(isinstance(style["hide_legend"],bool), "hide_legend must be a boolean.")
    if "reference_calibration" in style:
        require(isinstance(style["reference_calibration"],bool), "reference_calibration must be a boolean.")
        if style["reference_calibration"]:
            require(bool(style.get("slope_unit")) and not any(c in style["slope_unit"] for c in '\r\n;"'),
                    "Reference calibration needs a safe slope_unit.")
    keys(spec.get("export",{}), "formats width_px publication", "export")
    formats = spec.get("export",{}).get("formats",["png"])
    require(isinstance(formats,list) and set(formats) <= {"png","pdf","svg","tiff"}, "Unsupported export format.")
    require("png" in formats, "PNG is required alongside mandatory OPJU.")
    require("svg" not in formats, "Origin 2021 has no native SVG export (introduced in 2022b). Use PDF for vector output; do not mislabel PNG.")
    if spec.get("export",{}).get("publication"):
        require(bool(set(formats)&{"pdf","svg"}), "Publication figures require a requested vector format.")
    require(200 <= int(spec.get("export",{}).get("width_px",1800)) <= 20000, "Invalid PNG width.")
    ids = set()
    for d in spec["data"]:
        keys(d,"id source sheet header delimiter encoding inline units missing sort_by object","data")
        require(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_-]{0,63}",d.get("id","")) and d["id"] not in ids, "Data ids must be safe and unique.")
        ids.add(d["id"])
        require(d.get("object","worksheet") == "worksheet", "Matrix/XYZ fields need a custom script.")
        require(("source" in d) != ("inline" in d), "Provide exactly one source or inline data.")
        require(d.get("missing","error") in ("error","drop"), "Missing policy: error or explicit drop.")
    layers=set()
    for layer in spec["layers"]:
        keys(layer,"id index position","layer")
        require(layer.get("id") and layer["id"] not in layers, "Layer ids must be unique.")
        layers.add(layer["id"])
        if "position" in layer:
            pos=layer["position"]
            require(len(pos)==4 and all(math.isfinite(float(v)) for v in pos), "position is [left,top,width,height] in percent.")
            require(min(pos)>=0 and min(pos[2:])>0 and pos[0]+pos[2]<=100 and pos[1]+pos[3]<=100,"Layer outside page.")
    require([layer.get("index",i) for i,layer in enumerate(spec["layers"])]==list(range(len(layers))),"Layer indices must be sequential from zero.")
    plotids=set()
    for plot in spec["plots"]:
        keys(plot,"id layer data type x y y_error x_error color","plot")
        require(plot.get("id") and plot["id"] not in plotids,"Plot ids must be unique.")
        plotids.add(plot["id"])
        require(plot.get("layer") in layers and plot.get("data") in ids,"Unresolved plot data/layer.")
        require(plot.get("type") in ("line","scatter","line_symbol","bar"),"Plot type needs a custom Origin script.")
        require("x" in plot and "y" in plot and plot["x"]!=plot["y"],"Explicit distinct X and Y columns required.")
        if "color" in plot: require(re.fullmatch(r"#[0-9a-fA-F]{6}",plot["color"]),"Invalid plot color.")
    for lid,axes in spec.get("axes",{}).items():
        require(lid in layers,"Unknown axis layer.")
        keys(axes,"x y","axes")
        for axis in axes.values():
            keys(axis,"title scale limits","axis")
            require(axis.get("scale","linear") in ("linear","log10","ln","log2"),"Unsupported axis scale.")
            require(not any(c in axis.get("title","") for c in '\r\n;"'), "Axis title contains script delimiters; use custom label API.")
            if "limits" in axis:
                lim=axis["limits"]
                require(len(lim)==2 and all(math.isfinite(float(v)) for v in lim) and lim[0]<lim[1],"Invalid axis limits.")
                if axis.get("scale","linear")!="linear": require(lim[0]>0,"Log axis limits must be positive.")
    for a in spec.get("annotations",[]):
        keys(a,"layer text x y","annotation")
        require(a.get("layer") in layers,"Unknown annotation layer.")
        require("text" in a and "x" in a and "y" in a,"Text annotations need text,x,y.")
    for fit in spec.get("analysis",[]):
        keys(fit,"type plot annotate","analysis")
        require(fit.get("type")=="linear" and fit.get("plot") in plotids,"Advanced analysis needs a reviewed custom script.")
    template=spec.get("page",{}).get("template")
    if template:
        candidate=P(template)
        if candidate.suffix:
            if not candidate.is_absolute(): candidate=workspace/"templates"/candidate
            require(candidate.is_file() and candidate.suffix.lower() in (".otp",".otpu"),"Template not found or wrong type.")
    return spec

def prepare(spec, workspace):
    import pandas as pd
    import numpy as np
    frames, audits = {}, {}
    for item in spec["data"]:
        if "source" in item:
            source=P(item["source"])
            if not source.is_absolute(): source=workspace/source
            df,source=read_table(source,item.get("sheet",0),item.get("header",0),item.get("delimiter"),item.get("encoding","utf-8-sig"))
            audit={"source":str(source),"sha256":sha(source)}
        else:
            df=pd.DataFrame(item["inline"])
            df.columns=[str(c) for c in df.columns]
            audit={"source":"synthetic inline values","sha256":None}
        require(not df.empty,"Data has no rows.")
        used=[]
        for p in spec["plots"]:
            if p["data"]==item["id"]:
                for role in ("x","y","y_error","x_error"):
                    if role in p and p[role] not in used: used.append(p[role])
        require(used, "Unused data source; remove it or map it.")
        for col in used:
            require(col in df.columns,f"Column {col!r} is absent; available: {list(df.columns)}")
            converted=pd.to_numeric(df[col],errors="coerce")
            require(not (converted.isna() & df[col].notna()).any(), f"Non-numeric values in {col}; inspect units/categories before plotting.")
            df[col]=converted
        bad=~np.isfinite(df[used].to_numpy(dtype=float)).all(axis=1)
        audit.update(original_rows=len(df), excluded_rows_1based=(np.where(bad)[0]+1).tolist(), missing_policy=item.get("missing","error"))
        if bad.any():
            mappings={tuple(p.get(k) for k in ("x","y","x_error","y_error"))
                      for p in spec["plots"] if p["data"]==item["id"]}
            require(len(mappings)==1 or item.get("missing","error")!="drop",
                    "Shared multi-series data contain missing values. Split into separate DataSpec ids per series "
                    "before drop, or use a custom script with an explicit common-row policy.")
            require(item.get("missing","error")=="drop","Missing/non-finite values require an explicit data policy.")
            df=df.loc[~bad].copy()
        if item.get("sort_by"):
            require(all(c in df.columns for c in item["sort_by"]),"Unknown sort column.")
            df=df.sort_values(item["sort_by"],kind="stable")
        require(not df.empty,"No usable rows remain.")
        audit.update(plotted_rows=len(df),sort_by=item.get("sort_by",[]),normalized=False)
        frames[item["id"]]=df.reset_index(drop=True); audits[item["id"]]=audit
    for plot in spec["plots"]:
        df=frames[plot["data"]]
        for role in ("y_error","x_error"):
            if role in plot: require((df[plot[role]]>=0).all(),"Error magnitudes must be nonnegative.")
        for axis in ("x","y"):
            scale=spec.get("axes",{}).get(plot["layer"],{}).get(axis,{}).get("scale","linear")
            if scale!="linear": require((df[plot[axis]]>0).all(),"Nonpositive data on log axis: decide policy explicitly.")
    for fit in spec.get("analysis",[]):
        plot=next(p for p in spec["plots"] if p["id"]==fit["plot"])
        vals=frames[plot["data"]][plot["x"]]
        require(len(vals)>=3 and vals.nunique()>1,"Linear fit needs at least 3 rows and distinct X values.")
    return frames,audits

def file_check(path, started):
    path=P(path)
    require(path.is_file() and path.stat().st_size>0,f"Missing or empty output: {path}")
    require(path.stat().st_mtime>=started-2,f"Stale output: {path}")
    with path.open("rb") as f: header=f.read(4096)
    ext=path.suffix.lower()
    if ext==".png": require(header.startswith(b"\x89PNG\r\n\x1a\n"),"Invalid PNG.")
    elif ext==".pdf": require(header.startswith(b"%PDF"),"Invalid PDF.")
    elif ext==".svg": require(b"<svg" in header,"Invalid SVG.")
    elif ext in (".tiff",".tif"): require(header[:4] in (b"II*\x00",b"MM\x00*"),"Invalid TIFF.")
    result={"path":str(path),"bytes":path.stat().st_size,"mtime":path.stat().st_mtime,"sha256":sha(path)}
    if ext==".png": result["pixels"]=list(struct.unpack(">II",header[16:24]))
    return result

def execute(spec, config, frames, audits):
    import numpy as np
    workspace=P(config["WORKSPACE"]).resolve()
    run_id=spec["id"]+"_"+dt.datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    logdir=workspace/"logs"/run_id
    logdir.mkdir(parents=True,exist_ok=False)
    report={"run_id":run_id,"status":"running","environment":environment(config),"data":audits,"checks":{}}
    started=time.time()
    write_json(logdir/"spec.json",spec)
    lock=(workspace/"logs"/"origin_session.lock").open("a+b")
    lock.seek(0); lock.write(b"0"); lock.flush(); lock.seek(0)
    try:
        msvcrt.locking(lock.fileno(),msvcrt.LK_NBLCK,1)
    except OSError:
        lock.close()
        raise RuntimeError("Another Origin task owns this workspace session lock.")
    with (logdir/"run.log").open("w",encoding="utf-8") as logfile:
        def log(msg):
            line=dt.datetime.now().isoformat(timespec="seconds")+" "+msg
            print(line,flush=True); logfile.write(line+"\n"); logfile.flush()
        op=None
        try:
            import originpro as op
            before=origin_processes(); report["processes_before"]=before
            log("Connecting to a dedicated Origin Application session")
            op.set_show(bool(spec.get("runtime",{}).get("show_origin",True)))
            report["origin_version"]=op.lt_float("@V"); report["origin_bits"]=op.lt_float("@VB")
            require(9.8<=report["origin_version"]<9.85 and report["origin_bits"]==64,"Wrong Origin version/bitness; expected Origin 2021 64-bit.")
            after=origin_processes(); report["processes_during"]=after
            report["checks"]["origin_started"]=bool(after)
            require(after,"Origin64 process was not observed.")
            op.new()
            sheets={}; books={}
            for item in spec["data"]:
                df=frames[item["id"]]; w=op.new_sheet("w",lname=item["id"]); w.from_df(df)
                sheets[item["id"]]=w; books[item["id"]]=w.get_book().name
                for col,unit in item.get("units",{}).items():
                    require(col in df.columns,"Unit refers to unknown column.")
                    w.set_label(list(df.columns).index(col),unit,"U")
                for plot in spec["plots"]:
                    if plot["data"]!=item["id"]: continue
                    for role,des in (("x","X"),("y","Y"),("x_error","XErr"),("y_error","YErr")):
                        if role in plot:
                            idx=list(df.columns).index(plot[role])
                            # from_list axis accepts X,Y,E (Y error), XErr via cols_axis later if needed.
                            if role in ("x","y"): w.cols_axis(des.lower(),c1=idx,c2=idx)
                for plot in spec["plots"]:
                    if plot["data"]==item["id"]:
                        for role in ("x","y","x_error","y_error"):
                            if role in plot:
                                col=list(df.columns).index(plot[role])
                                require(np.allclose(w.to_list(col),df[plot[role]].to_numpy(dtype=float),rtol=1e-12,atol=1e-12),"Worksheet data mismatch.")
                archive=workspace/"output"/"data"/run_id/(item["id"]+".csv")
                archive.parent.mkdir(parents=True,exist_ok=True); df.to_csv(archive,index=False,encoding="utf-8-sig")
                log(f"Worksheet {books[item['id']]}: {len(df)} rows verified")
            template=spec.get("page",{}).get("template")
            if template and P(template).suffix and not P(template).is_absolute():
                template=str(workspace/"templates"/template)
            graph=op.new_graph(template=template or "Origin",lname=spec["id"])
            layers={}
            for i,ls in enumerate(spec["layers"]):
                while len(graph)<=i: graph.add_layer()
                layer=graph[i]; layers[ls["id"]]=layer
                if "position" in ls:
                    layer.set_int("unit",1)
                    for k,v in zip(("left","top","width","height"),ls["position"]): layer.set_float(k,float(v))
            style=spec.get("style",{})
            colors=style.get("colors",["#2463A6","#C65A32","#258875","#7B62A3"])
            created={}
            expected={lid:0 for lid in layers}
            for i,ps in enumerate(spec["plots"]):
                df=frames[ps["data"]]; names=list(df.columns); layer=layers[ps["layer"]]
                kwargs={r:names.index(ps[k]) for k,r in (("x_error","colxerr"),("y_error","colyerr")) if k in ps}
                plot=layer.add_plot(sheets[ps["data"]],colx=names.index(ps["x"]),coly=names.index(ps["y"]),
                                   type={"line":"l","scatter":"s","line_symbol":"y","bar":"c"}[ps["type"]],**kwargs)
                require(plot is not None,"Origin failed to create a plot.")
                # Preserve template visuals unless a style override was specified.
                if not template or style or ps.get("color"):
                    plot.color=ps.get("color",colors[i%len(colors)])
                    plot.symbol_size=float(style.get("symbol_size",6))
                    plot.set_cmd("-w "+str(500*float(style.get("line_width",1.5))))
                created[ps["id"]]=plot; expected[ps["layer"]]+=1
            fit_results=[]
            for fit in spec.get("analysis",[]):
                ps=next(p for p in spec["plots"] if p["id"]==fit["plot"])
                df=frames[ps["data"]]; names=list(df.columns)
                # Origin 2021: high-level LinearFit result XML can be truncated; use native fitLR scalars.
                w=sheets[ps["data"]]; xr=names.index(ps["x"])+1; yr=names.index(ps["y"])+1
                op.lt_exec("double __oa=NANUM; double __ob=NANUM; double __oae=NANUM; double __obe=NANUM; int __on=-1;")
                op.lt_exec(f"fitLR iy:={w.lt_range(False)}!({xr},{yr}) a:=__oa b:=__ob aerr:=__oae berr:=__obe pts:=__on;")
                intercept=op.lt_float("__oa"); slope=op.lt_float("__ob")
                require(math.isfinite(slope) and math.isfinite(intercept),"Origin fitLR returned invalid parameters.")
                require(int(op.lt_float("__on"))==len(df),"Origin fitLR used an unexpected number of rows.")
                result={"intercept":intercept,"slope":slope,"points":int(op.lt_float("__on")),
                        "intercept_error":op.lt_float("__oae"),"slope_error":op.lt_float("__obe")}
                for key,val in list(result.items()):
                    if isinstance(val,float) and not math.isfinite(val): result[key]=None
                x=df[ps["x"]].to_numpy(dtype=float); y=df[ps["y"]].to_numpy(dtype=float)
                residual=y-(slope*x+intercept); sst=float(np.sum((y-y.mean())**2))
                r2=1-float(np.sum(residual**2))/sst if sst>0 else None
                fitw=op.new_sheet("w",lname="LinearFit_"+ps["id"])
                fitx=np.linspace(float(x.min()),float(x.max()),101); fity=slope*fitx+intercept
                fitw.from_list(0,fitx.tolist(),lname=ps["x"],axis="X")
                fitw.from_list(1,fity.tolist(),lname="Linear fit",axis="Y")
                fp=layers[ps["layer"]].add_plot(fitw,1,0,type="l"); require(fp is not None,"Fit curve missing.")
                fp.color=style.get("fit_line_color","#C65A32")
                fp.set_cmd("-w 500" if style.get("reference_calibration") else "-w 750")
                expected[ps["layer"]]+=1
                fr={"plot":ps["id"],"engine":"Origin fitLR (native LabTalk)","slope":slope,"intercept":intercept,
                    "r_squared":r2,"r_squared_definition":"1-SSE/SST (unweighted, from Origin coefficients)",
                    "weighting":"unweighted","native_result":result}
                fit_results.append(fr)
                log(f"Origin linear fit: slope={slope:.8g}, intercept={intercept:.8g}, R2={r2}")
            for lid,layer in layers.items():
                require(expected[lid]>0,"Every layer must contain a requested plot.")
                layer.rescale()
                for axis,axspec in spec.get("axes",{}).get(lid,{}).items():
                    ax=layer.axis(axis)
                    if "title" in axspec: ax.title=axspec["title"]
                    if "scale" in axspec: ax.scale=axspec["scale"]
                    if "limits" in axspec: ax.limits=axspec["limits"]
                if not template or style:
                    font=style.get("font","Arial"); fs=float(style.get("font_size",16))
                    layer.lt_exec(f'layer.x.label.font=font({font}); layer.y.label.font=font({font}); '
                                  f'layer.x.label.fsize={fs}; layer.y.label.fsize={fs}; '
                                  f'layer.x.thickness={2.1 if style.get("reference_calibration") else 1.2}; '
                                  f'layer.y.thickness={2.1 if style.get("reference_calibration") else 1.2}; '
                                  'layer.x.grid=0; layer.y.grid=0; layer.x.minorGrid=0; layer.y.minorGrid=0;')
                    if style.get("reference_calibration"):
                        layer.lt_exec('layer.x.opposite=1; layer.y.opposite=1; '
                                      'layer.x2.ticks=0; layer.y2.ticks=0; '
                                      'layer.x.ticks=2; layer.y.ticks=2; '
                                      'layer.x.minorTicks=0; layer.y.minorTicks=0; '
                                      'layer.x.ticklength=8; layer.y.ticklength=8;')
                    for label_name in ("xb","yl"):
                        label=layer.label(label_name)
                        if label:
                            label.set_int("font",int(op.lt_float(f"font({font})")))
                            label.set_float("fsize",fs+2)
                if expected[lid]==1 or style.get("hide_legend",False):
                    layer.remove_label("legend")
                elif not template:
                    layer.lt_exec("legend -r; legend.background=0; legend.fsize=12; legend.font=font(Arial);")
                require(len(layer.plot_list())>=expected[lid], "Graph layer has fewer plots than expected.")
            annotation_checks=[]
            for fit,fr in zip(spec.get("analysis",[]),fit_results):
                if fit.get("annotate",True):
                    ps=next(p for p in spec["plots"] if p["id"]==fit["plot"])
                    layer=layers[ps["layer"]]; xmin,xmax,_=layer.xlim; ymin,ymax,_=layer.ylim
                    if style.get("reference_calibration"):
                        x=xmin+.55*(xmax-xmin)
                        lines=[(f"K={fr['slope']:.6g}{style['slope_unit']}", ymin+.34*(ymax-ymin))]
                        if fr["r_squared"] is not None:
                            lines.append((f"R² = {fr['r_squared']:.5f}", ymin+.27*(ymax-ymin)))
                        for text,y in lines:
                            label=layer.add_label(text,x,y)
                            require(label is not None,"Reference calibration annotation failed.")
                            label.set_float("x1",x); label.set_float("y1",y)
                            label.set_float("fsize",17)
                            label.set_int("font",int(op.lt_float(f"font({style.get('font','Arial')})")))
                            annotation_checks.append((ps["layer"],label.name,text))
                    else:
                        text=f"Y = {fr['slope']:.5g} X {fr['intercept']:+.5g}"
                        if fr["r_squared"] is not None: text+=f"\nR\\+(2) = {fr['r_squared']:.5f}"
                        label=layer.add_label(text,xmin+.05*(xmax-xmin),ymax-.08*(ymax-ymin))
                        if label: label.set_float("fsize",12)
            for ann in spec.get("annotations",[]):
                label=layers[ann["layer"]].add_label(ann["text"],ann["x"],ann["y"])
                require(label is not None,"Annotation creation failed.")
                # add_label treats zero as absent; explicitly set coordinates afterward.
                label.set_float("x1",float(ann["x"])); label.set_float("y1",float(ann["y"]))
            graph_name=graph.name
            report["checks"].update(workbooks_created=bool(books),worksheet_values_verified=True,
                rows={k:len(v) for k,v in frames.items()},graph_page=graph_name,
                plots_per_layer={k:len(v.plot_list()) for k,v in layers.items()})
            report["analysis"]=fit_results
            project=workspace/"output"/"opju"/(run_id+".opju")
            require(op.save(str(project)),"Origin OPJU save failed before export.")
            report["saved_project_before_export"]=str(project)
            files=[]
            for ext in spec.get("export",{}).get("formats",["png"]):
                export_type="tif" if ext=="tiff" else ext
                path=workspace/"output"/"figures"/(run_id+"."+export_type)
                width=int(spec.get("export",{}).get("width_px",1800)) if ext in ("png","tiff") else 0
                result=graph.save_fig(str(path),type=export_type,width=width,replace=False)
                require(result,f"Origin export failed for {ext}.")
                files.append(file_check(path,started)); log(f"Exported {path.name}")
            project=workspace/"output"/"opju"/(run_id+".opju")
            require(op.save(str(project)),"Origin OPJU save failed."); files.append(file_check(project,started))
            seconds=float(spec.get("runtime",{}).get("show_seconds",3))
            require(0<=seconds<=60,"show_seconds must be in 0..60.")
            time.sleep(seconds)
            # Release wrappers before reloading; look up worksheets explicitly, not by active graph.
            sheets.clear(); created.clear(); layers.clear()
            graph=layer=plot=None
            require(op.open(str(project)),"Saved OPJU could not be reopened.")
            for did,book in books.items():
                w=op.find_sheet("w",book); require(w is not None,"Saved workbook absent.")
                df=frames[did]
                for ps in spec["plots"]:
                    if ps["data"]!=did: continue
                    for role in ("x","y","x_error","y_error"):
                        if role in ps:
                            col=list(df.columns).index(ps[role]); values=w.to_list(col)
                            require(len(values)==len(df) and np.allclose(values,df[ps[role]],rtol=1e-12,atol=1e-12),"Persisted worksheet rows/values differ.")
            g=op.find_graph(graph_name); require(g is not None,"Saved graph page absent.")
            for i,ls in enumerate(spec["layers"]):
                require(len(g[i].plot_list())>=expected[ls["id"]],"Saved plot absent.")
                for layer_id,name,text in annotation_checks:
                    if layer_id == ls["id"]:
                        saved_label=g[i].label(name)
                        require(saved_label is not None and saved_label.text==text,
                                "Saved K or R² annotation missing from Origin graph.")
            report["checks"]["reference_annotations_reopened"]=len(annotation_checks)
            for did,audit in audits.items():
                if audit["sha256"]: require(sha(audit["source"])==audit["sha256"],"Source data changed during job.")
            report["files"]=files; report["checks"].update(opju_reopened=True,source_hashes_unchanged=True)
            report["status"]="success"; log("QA passed: Origin objects, data, plots, files, timestamps and reopened OPJU")
        except BaseException as exc:
            report["status"]="failed"; report["error"]=str(exc)
            logfile.write(traceback.format_exc()); logfile.flush()
            raise
        finally:
            if op is not None:
                try:
                    op.exit(); report["origin_exit"]="ok"; log("Origin exit OK")
                except Exception as exc:
                    report["origin_exit"]=str(exc); report["status"]="failed"
            report["finished_at"]=dt.datetime.now().isoformat()
            write_json(logdir/"qa.json",report)
            write_json(workspace/"logs"/"latest.json",{"run_id":run_id,"qa":str(logdir/"qa.json"),"status":report["status"]})
            lock.seek(0); msvcrt.locking(lock.fileno(),msvcrt.LK_UNLCK,1); lock.close()
    require(report["status"]=="success","Job or Origin shutdown failed; inspect qa.json.")
    return report

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config",required=True,type=P)
    sub=parser.add_subparsers(dest="command",required=True)
    sub.add_parser("probe")
    inspect=sub.add_parser("inspect"); inspect.add_argument("file",type=P); inspect.add_argument("--sheet",default="0")
    for cmd in ("preview","run"):
        p=sub.add_parser(cmd); p.add_argument("spec",type=P)
    args=parser.parse_args(); config=read_json(args.config); environment(config)
    workspace=P(config["WORKSPACE"]).resolve()
    if args.command=="probe":
        print(json.dumps(environment(config),indent=2)); return
    if args.command=="inspect":
        sheet=int(args.sheet) if args.sheet.isdecimal() else args.sheet
        print(json.dumps(inspect_data(args.file,sheet),ensure_ascii=False,indent=2)); return
    spec=validate(read_json(args.spec),workspace); frames,audits=prepare(spec,workspace)
    if args.command=="preview" or spec["mode"]=="preview":
        print(json.dumps({"mode":"preview","origin_started":False,"plan":spec,"data_audit":audits},ensure_ascii=False,indent=2))
        return
    execute(spec,config,frames,audits)

if __name__=="__main__":
    main()
